#!/usr/bin/env python3
"""
Gera um script SQL (para colar no SQL Editor do Supabase) que cadastra peças e preços
a partir de listas de fornecedores em Excel.

Uso:
  python3 scripts/gerar_importacao.py LISTA.xlsx --data 2026-09-30 --saida imports/lista.sql

Lista "Vibe / Marítimo" (uma aba, 3 colunas de modelos com preço ao lado):
  - nome em vermelho (ou com "VIBE" no nome) -> fornecedor Vibe
  - nome em preto                           -> fornecedor Marítimo
  - itens sem preço são ignorados.
Peças com o mesmo nome (sem a palavra VIBE, sem acentos/pontuação) viram UMA peça com um preço por fornecedor.
O script é reexecutável: preços diferentes atualizam a oferta (e geram histórico).
"""
import argparse, re, sys, unicodedata
from collections import defaultdict
from datetime import date as calendar_date
from decimal import Decimal, InvalidOperation
import openpyxl

BRANDS = {
    "SAMSUNG": "Samsung", "SANSUNG": "Samsung", "IPHONE": "Apple", "MOTOROLA": "Motorola", "LG": "LG",
    "XIAOME": "Xiaomi", "XIAOMI": "Xiaomi", "NOKIA": "Nokia", "INFINIX": "Infinix", "REALME": "Realme",
    "OPPO": "Oppo", "ITEL": "Itel", "HUAWEI": "Huawei", "ASUS": "Asus", "ZENFONE": "Asus",
}
NO_BRAND_SECTIONS = {"TOUCH", "LCD", "TABLET"}


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def clean_name(raw: str) -> str:
    s = re.sub(r"\bVIBE\b", " ", str(raw).upper())
    s = re.sub(r"\bINCEL\b", "INCELL", s).replace("C0RE", "CORE")
    return re.sub(r"\s+", " ", s).strip()


def display_name(name: str, brand: str | None) -> str:
    """"Tela " + marca (se o nome ainda não a traz) + nome da lista."""
    if brand == "Apple":
        rest = name[len("IPHONE"):].strip() if name.startswith("IPHONE") else name
        return f"Tela iPhone {rest}"
    if brand and not name.startswith(brand.upper()):
        return f"Tela {brand} {name}"
    return f"Tela {name}"


def technology(name: str) -> str | None:
    n = f" {name} "
    if " SOFT OLED " in n or " SOFT " in n and "OLED" in n:
        return "Soft OLED"
    for key, label in ((" OLED ", "OLED"), (" INCELL ", "Incell"), (" LCD ", "LCD"), (" AMOLED ", "AMOLED")):
        if key in n:
            return label
    return None


def quality(name: str) -> str | None:
    return "Original" if re.search(r"\b(ORI|ORIGINAL)\b", name) else None


def is_red(cell) -> bool:
    c = cell.font.color
    return c is not None and c.type == "rgb" and str(c.rgb).upper().endswith("FF0000")


def parse_vibe_maritimo(path: str):
    ws = openpyxl.load_workbook(path).active
    current = {1: "SAMSUNG", 4: "IPHONE", 7: "MOTOROLA"}  # coluna do nome -> seção atual
    rows, ignored = [], 0
    for row in ws.iter_rows(min_row=2):
        for ci in (1, 4, 7):
            cell, price = row[ci], row[ci + 1].value
            if cell.value is None or not str(cell.value).strip():
                continue
            raw = str(cell.value).strip()
            if isinstance(price, str) and price.strip().upper().startswith("PRE"):
                current[ci] = raw.upper()          # cabeçalho de seção (LG, XIAOME, ...)
                continue
            if price is None and raw.upper() in BRANDS:
                current[ci] = raw.upper()          # sub-cabeçalho sem preço (OPPO, ITEL)
                continue
            if not isinstance(price, (int, float)) or isinstance(price, bool) or price <= 0:
                ignored += 1
                continue
            section = current[ci]
            brand = None if section in NO_BRAND_SECTIONS else BRANDS.get(section)
            supplier = "Vibe" if (is_red(cell) or re.search(r"\bVIBE\b", raw.upper())) else "Marítimo"
            name = clean_name(raw)
            rows.append({"name": display_name(name, brand), "brand": brand, "tech": technology(name),
                         "quality": quality(name), "supplier": supplier, "price": float(price),
                         "src": f"linha {cell.row}"})
    return rows, ignored


# ----------------------------------------------------------------------------------------
# Lista World Mobile: blocos de 3 colunas (MODELOS | SEM ARO | COM ARO), preços como "R$ 45,00".
# Um modelo pode agrupar vários aparelhos ("A02S-A03-A04E"): cada aparelho vira uma peça,
# com o mesmo preço, para casar com as listas dos outros fornecedores.
# ----------------------------------------------------------------------------------------
WM_TRIPLETS = [(2, 3, 4), (5, 6, 7), (8, 9, 10), (11, 12, 13), (14, 15, 16), (17, 18, 19)]
# cabeçalho (linha 5) -> estado inicial: (marca, sufixo, divide por "-")
WM_START = {
    2: ("Samsung", "", True), 5: ("Samsung", "", True), 8: ("Motorola", "", True),
    11: ("LG", "", True), 14: ("Asus", "", True), 17: ("Xiaomi", "", True),
}
# subcabeçalhos (texto sem preço) -> (marca, sufixo); None mantém o valor atual
WM_HEADERS = [
    # linhas "PRIME / WEFIX / NACIONAL / DIAGNOSTIC" são outras qualidades, com nomes e preços que
    # não casam com as demais listas: ficam de fora (SKIP vale para o resto da coluna).
    ("WEFIX", (None, "SKIP")), ("SAMSUNG NACIONAL", ("Samsung", "SKIP")), ("FRONTAL NACIONAL", (None, "SKIP")),
    ("REALME LINHA PRIME", ("Realme", "SKIP")),
    ("MOTOROLA OLED", ("Motorola", "")), ("FRONTAL XIAOMI OLED", ("Xiaomi", "OLED")),
    ("IPHONE INCELL", ("Apple", "INCELL")), ("IPHONE OLED", ("Apple", "OLED")),
    ("REALME", ("Realme", "")), ("HUAWEI", ("Huawei", "")), ("OPPO", ("Oppo", "")), ("HONOR", ("Honor", "")),
    ("INFINIX", ("Infinix", "")), ("TOUCH TABLET", (None, "TABLET")),
]


def parse_price(v):
    """Preço completo e não negativo; nunca extrai um número de texto inválido."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, str):
        s = re.sub(r"^R\$\s*", "", v.strip(), flags=re.I)
        if "," in s:
            if not re.fullmatch(r"(?:\d+|\d{1,3}(?:\.\d{3})+),\d{1,2}", s):
                return None
            s = s.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"\d{1,3}(?:\.\d{3})+", s):
            s = s.replace(".", "")
        elif not re.fullmatch(r"\d+(?:\.\d{1,2})?", s):
            return None
    elif isinstance(v, (int, float, Decimal)):
        s = str(v)
    else:
        return None
    try:
        n = Decimal(s)
        if not n.is_finite() or n < 0 or n >= Decimal("10000000000"):
            return None
        return n.quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def wm_price(v):
    n = parse_price(v)
    return float(n) if n is not None and n > 0 else None


def import_date(value: str) -> str:
    """Valida antes de interpolar datas no SQL (inclusive as datas por linha)."""
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise ValueError("Data inválida: use AAAA-MM-DD")
    return calendar_date.fromisoformat(value).isoformat()


def wm_models(raw: str, brand: str | None) -> list[str]:
    s = re.sub(r"\s+", " ", raw.upper()).strip()
    s = re.sub(r"\b(PROMOCAO|PROMOÇÃO)\b", "", s)
    s = re.sub(r"\bINCEEL\b", "INCELL", s)
    s = re.sub(r"\bINCEL\b", "INCELL", s)
    s = re.sub(r"^(MT|FRONTAL)\s+", "", s)
    s = re.sub(r"\s+", " ", s).strip(" -/")
    trailing = ""
    m = re.search(r"\b(OLED|INCELL)\b", s)
    if m and brand != "Apple":
        trailing = m.group(1)                      # "A16 OLED/A17/A26": a tecnologia vale para todos
        s = re.sub(r"\s+", " ", re.sub(rf"\b{trailing}\b", " ", s)).strip(" -/")
    seps = r"/" if brand == "Apple" else r"[-/]"
    parts = [p.strip() for p in re.split(seps, s) if p.strip()]
    out, prev = [], ""
    for p in parts:
        if re.fullmatch(r"[45]G", p) and prev:           # "A16 4G/5G" -> A16 4G, A16 5G
            p = re.sub(r"\s*[45]G$", "", prev) + " " + p
        if not re.search(r"\d", p):            # sem número = não é um código de modelo ("OLED", "F. DIGITAL")
            continue
        if trailing and not p.endswith(trailing):
            p = f"{p} {trailing}"
        out.append(p)
        prev = p
    return out


def parse_world_mobile(path: str):
    ws = openpyxl.load_workbook(path, data_only=True).active
    state = {m: list(v[:2]) for m, v in WM_START.items()}
    rows, skipped = [], defaultdict(int)
    for r in range(7, ws.max_row + 1):
        for m, s_col, c_col in WM_TRIPLETS:
            raw = ws.cell(r, m).value
            if raw is None or not str(raw).strip():
                continue
            raw = str(raw).strip()
            sem, com = wm_price(ws.cell(r, s_col).value), wm_price(ws.cell(r, c_col).value)
            if sem is None and com is None:
                if state[m][1] == "SKIP":
                    continue
                for key, (b, suf) in WM_HEADERS:
                    if raw.upper() == key or raw.upper().startswith(key):
                        if b:
                            state[m][0] = b
                        state[m][1] = suf
                        break
                continue
            brand, suffix = state[m]
            if re.search(r"PRIME|DIAGNOSTIC|RETIRAD|NACIONAL|PROMO|\bT/C\b", raw.upper()):
                skipped["linhas PRIME/NACIONAL/WEFIX"] += 1
                continue
            if suffix in ("SKIP", "TABLET"):
                skipped["linhas PRIME/NACIONAL/WEFIX" if suffix == "SKIP" else "tablets"] += 1
                continue
            for model in wm_models(raw, brand):
                tech = technology(model)
                variants = [("SEM ARO", sem), ("COM ARO", com)]
                if brand == "Apple":                      # iPhone não tem "aro": um preço só
                    variants = [("", sem if sem is not None else com)]
                for var, price in variants:
                    if price is None:
                        continue
                    suf = suffix if suffix not in ("", "TABLET") and suffix not in model.split() else ""
                    name = " ".join(x for x in (model, suf, var) if x)
                    rows.append({"name": display_name(name, brand), "brand": brand, "tech": tech,
                                 "quality": quality(name), "supplier": "World Mobile", "price": price,
                                 "src": f"célula {openpyxl.utils.get_column_letter(m)}{r}: {raw}"})
    return rows, dict(skipped)


def q(v):
    return "null" if v is None else "'" + str(v).replace("'", "''") + "'"


def build_sql(rows, date: str, suppliers: list[str], conflict: str = "min"):
    date = import_date(date)
    # um preço por (peça, fornecedor). Em duplicidade com preços diferentes:
    #   conflict="min"     -> vale o menor
    #   conflict="exclude" -> a peça NÃO é importada para esse fornecedor (vai para o relatório)
    groups = defaultdict(list)
    for r in rows:
        r = dict(r)
        price = parse_price(r["price"])
        if price is None:
            raise ValueError(f"Preço inválido em {r['name']}: {r['price']!r}")
        r["price"] = price
        if r.get("date") is not None:
            r["date"] = import_date(r["date"])
        groups[(norm(r["name"]), r["supplier"])].append(r)
    final, dup = [], []
    for k, g in groups.items():
        if len(g) == 1:
            final.append(g[0])
            continue
        prices = {x["price"] for x in g}
        dup.append(g)
        if len(prices) == 1 or conflict == "min":
            final.append(min(g, key=lambda x: x["price"]))
    if not final:
        return "-- Nenhuma oferta válida para importar; nenhum dado foi alterado.\n", final, dup
    values = ",\n".join(
        f"  ({q(r['name'])}, {q(r.get('category', 'Tela'))}, {q(r['brand'])}, {q(r.get('model'))}, {q(r['tech'])}, {q(r['quality'])}, "
        f"{q(r.get('notes'))}, {q(r['supplier'])}, {q(r.get('code'))}, {q(r.get('offer_notes'))}, "
        f"{q(r.get('availability', 'in_stock'))}, {q(r.get('date'))}, {r['price']:.2f})"
        for r in final)
    sup_values = ", ".join(f"({q(s)})" for s in suppliers)
    sql = f"""-- =============================================================================
-- PointParts — importação de lista de preços ({date})
-- Fornecedores: {', '.join(suppliers)} | Ofertas: {len(final)} | Peças distintas: {len({norm(r['name']) for r in final})}
-- Cole no SQL Editor do Supabase e clique em Run. Pode ser executado de novo:
-- peças já existentes são reaproveitadas e preços diferentes atualizam a oferta (e geram histórico).
-- =============================================================================
begin;

-- evita que duas execuções simultâneas criem peças duplicadas
select pg_advisory_xact_lock(hashtext('pointparts_import'));

create temp table _imp (
  name text, category text, brand text, device_model text, technology text, quality text, notes text,
  supplier text, supplier_code text, offer_notes text, availability text, checked text, price numeric(12,2)
) on commit drop;

insert into _imp (name, category, brand, device_model, technology, quality, notes, supplier, supplier_code, offer_notes, availability, checked, price) values
{values};

-- fornecedores
insert into public.suppliers (name)
select v.n from (values {sup_values}) as v(n)
where not exists (select 1 from public.suppliers s where lower(s.name) = lower(v.n));

-- categorias novas (ex.: Carregador, Cabo) são criadas se ainda não existirem
insert into public.categories (name)
select distinct i.category from _imp i
where i.category is not null
  and not exists (select 1 from public.categories c where lower(c.name) = lower(i.category));

-- peças novas; as que já existem (mesmo nome, sem acento/pontuação) são reaproveitadas
insert into public.parts (name, category_id, brand, device_model, technology, quality, notes)
select distinct on (public.norm_text(i.name))
       i.name, (select id from public.categories where lower(name) = lower(i.category)), i.brand, i.device_model, i.technology, i.quality, i.notes
from _imp i
where not exists (select 1 from public.parts p where public.norm_text(p.name) = public.norm_text(i.name))
order by public.norm_text(i.name);

-- índice temporário nome normalizado -> peça (uma só leitura da tabela de peças)
create temp table _pn on commit drop as
select distinct on (n) id, n
from (select id, public.norm_text(name) as n, created_at from public.parts) x
order by n, created_at;

-- ofertas (um preço por fornecedor); data da consulta = data da lista
-- (o gatilho que atualiza a busca a cada oferta é pausado para ser rápido; a busca é recalculada no final)
alter table public.supplier_parts disable trigger supplier_parts_refresh_part;

insert into public.supplier_parts (part_id, supplier_id, price, availability, supplier_code, notes, checked_at)
select p.id, s.id, i.price, coalesce(i.availability, 'in_stock'), i.supplier_code, i.offer_notes,
       coalesce((i.checked || ' 12:00:00-03')::timestamptz, timestamptz '{date} 12:00:00-03')
from _imp i
join public.suppliers s on lower(s.name) = lower(i.supplier)
join _pn p on p.n = public.norm_text(i.name)
on conflict (part_id, supplier_id) do update
  set price = excluded.price, availability = excluded.availability,
      supplier_code = coalesce(excluded.supplier_code, public.supplier_parts.supplier_code),
      notes = coalesce(excluded.notes, public.supplier_parts.notes), checked_at = excluded.checked_at
  -- Lista antiga não substitui consulta/preço mais recente (nem gera histórico falso).
  where excluded.checked_at >= public.supplier_parts.checked_at;

alter table public.supplier_parts enable trigger supplier_parts_refresh_part;

-- recalcula o texto de busca só das peças que receberam código de fornecedor
update public.parts p set name = p.name
 where p.id in (select sp.part_id from public.supplier_parts sp
                  join public.suppliers s on s.id = sp.supplier_id
                  join _imp i on lower(s.name) = lower(i.supplier)
                 where i.supplier_code is not null and sp.supplier_code is not null);

commit;

select (select count(*) from public.parts)          as total_pecas,
       (select count(*) from public.supplier_parts) as total_ofertas,
       (select count(*) from public.suppliers)      as total_fornecedores;
"""
    return sql, final, dup


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--data", required=True, help="data da lista, AAAA-MM-DD")
    ap.add_argument("--saida", required=True)
    ap.add_argument("--lista", choices=["vibe_maritimo", "world_mobile", "json"], default="vibe_maritimo")
    ap.add_argument("--existentes", help="xlsx Vibe/Marítimo, para informar quais peças da World Mobile já existem")
    ap.add_argument("--relatorio", help="CSV com o resultado do casamento de nomes (World Mobile)")
    ap.add_argument("--max-linhas", type=int, default=0, help="divide a saída em vários arquivos (modo json)")
    a = ap.parse_args()
    if a.lista == "world_mobile":
        return main_wm(a)
    if a.lista == "json":
        return main_json(a)
    rows, ignored = parse_vibe_maritimo(a.xlsx)
    sql, final, dup = build_sql(rows, a.data, ["Vibe", "Marítimo"])
    open(a.saida, "w", encoding="utf-8").write(sql)
    by = defaultdict(set)
    for r in final:
        by[norm(r["name"])].add(r["supplier"])
    both = sum(1 for s in by.values() if len(s) > 1)
    print(f"ofertas: {len(final)} | peças: {len(by)} | em ambos os fornecedores: {both} | sem preço (ignorados): {ignored}")
    for s in ("Vibe", "Marítimo"):
        print(f"  {s}: {sum(1 for r in final if r['supplier'] == s)} ofertas")
    for g in dup:
        print(f"  duplicado ({g[0]['name']} / {g[0]['supplier']}): {sorted({x['price'] for x in g})} -> vale o menor")


def main_json(a):
    """Lista já estruturada em JSON: [{name, supplier, price, category, brand, quality, code, notes, availability, date}, ...]"""
    import json
    rows = json.load(open(a.xlsx, encoding="utf-8"))
    for r in rows:
        r.setdefault("brand", None); r.setdefault("tech", None); r.setdefault("quality", None)
    suppliers = sorted({r["supplier"] for r in rows})
    # em duplicidade com preços diferentes a peça NÃO é importada para esse fornecedor (conflito)
    _, final, dup = build_sql(rows, a.data, suppliers, conflict="exclude")
    conflicts = [g for g in dup if len({x["price"] for x in g}) > 1]
    n = getattr(a, "max_linhas", 0) or len(final) or 1
    chunks = [final[i:i + n] for i in range(0, len(final), n)]
    base = a.saida[:-4] if a.saida.endswith(".sql") else a.saida
    for i, chunk in enumerate(chunks, 1):
        out = a.saida if len(chunks) == 1 else f"{base}_parte{i}de{len(chunks)}.sql"
        sql, _, _ = build_sql(chunk, a.data, suppliers)
        open(out, "w", encoding="utf-8").write(sql)
        print("gerado:", out, f"({len(chunk)} ofertas)")
    print(f"ofertas: {len(final)} | peças: {len({norm(r['name']) for r in final})} | fornecedores: {', '.join(suppliers)}")
    print(f"duplicados com preços diferentes (não importados): {len(conflicts)}")
    for g in conflicts[:15]:
        print("  -", g[0]["name"], sorted({x["price"] for x in g}))


def main_wm(a):
    import csv
    rows, skipped = parse_world_mobile(a.xlsx)
    sql, final, dup = build_sql(rows, a.data, ["World Mobile"], conflict="exclude")
    open(a.saida, "w", encoding="utf-8").write(sql)
    existing = set()
    if a.existentes:
        base, _ = parse_vibe_maritimo(a.existentes)
        existing = {norm(r["name"]) for r in base}
    matched = [r for r in final if norm(r["name"]) in existing]
    print(f"ofertas World Mobile: {len(final)} | casam com peças já existentes: {len(matched)} | peças novas: {len(final) - len(matched)}")
    if skipped:
        print("ignorados:", skipped)
    conflicts = [g for g in dup if len({x["price"] for x in g}) > 1]
    print(f"conflitos (mesma peça com preços diferentes na planilha, NÃO importados): {len(conflicts)}")
    if a.relatorio:
        with open(a.relatorio, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["situacao", "peca", "preco", "origem_na_planilha"])
            for r in sorted(final, key=lambda x: (norm(x["name"]) in existing, x["name"])):
                w.writerow(["JA EXISTIA" if norm(r["name"]) in existing else "NOVA", r["name"], f"{r['price']:.2f}".replace(".", ","), r["src"]])
            for g in conflicts:
                for x in g:
                    w.writerow(["CONFLITO - NAO IMPORTADA", x["name"], f"{x['price']:.2f}".replace(".", ","), x["src"]])


if __name__ == "__main__":
    main()
